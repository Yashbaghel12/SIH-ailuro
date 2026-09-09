"""
Generate a Gemma fine-tuning dataset from complaint case metadata.
--------------------------------------------------------------------
Takes SIH_Ailuro_Synthetic_Atrocity_Complaints.xlsx (case-level
metadata: delay_days, case_status, incident_type, etc.) and, for
each case, synthesizes a plausible victim check-in (text + derived
signals), then scores it with the rule-based engine to produce a
weak label (severity + indicators).

This is the "bootstrap from rule engine" approach: the labels
aren't hand-annotated by a clinician, but they're grounded in
realistic case attributes instead of being randomly assigned.
Before using this to fine-tune, have a domain expert (counselor/
NGO staff) spot-check and correct a sample — see README.

Usage:
    python generate_finetune_dataset.py \
        --input SIH_Ailuro_Synthetic_Atrocity_Complaints.xlsx \
        --output dataset.jsonl \
        --sample-size 400 --seed 42
"""

import argparse
import json
import random
import sys
import os

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from engine.distress_engine import CheckIn, DistressEngine

STATUS_PHRASES = {
    "Referred to another authority": "they keep referring my case to another department",
    "Under investigation": "the investigation is still going on",
    "Chargesheet filed": "the chargesheet has been filed but I haven't heard anything since",
    "Preliminary inquiry": "they're still doing a preliminary inquiry",
    "Closed - insufficient evidence": "they closed my case saying there wasn't enough evidence",
    "Pending review": "my case is just sitting pending review",
}

REASON_CLAUSES = {
    "Fear of retaliation": "I'm scared of what might happen if I push further",
    "Social pressure": "people in my community have been pressuring me to drop it",
    "Concern about family safety": "I'm worried about my family's safety because of this",
    "Financial constraints": "I can't afford to keep going back and forth for this",
}

INJURY_CLAUSES = [
    "I'm still recovering from what happened",
    "the injury still bothers me sometimes",
]

THREAT_CLAUSES = [
    "I've also been getting threatened to withdraw the complaint",
    "someone from the other side keeps following me",
    "I feel unsafe going out alone now",
]

MOOD_CLAUSES = {
    1: ["I feel hopeless about all this", "I don't know how much longer I can keep going",
        "I can't sleep properly anymore"],
    2: ["I'm finding it hard to stay hopeful", "some days I just feel exhausted by it all"],
    3: ["I'm managing but it's tough some days"],
    4: ["I'm doing okay, trying to stay hopeful"],
    5: ["I'm feeling alright, staying positive"],
}


def delay_phrase(delay_days: int) -> str:
    if delay_days < 15:
        return "I filed my complaint recently"
    elif delay_days < 60:
        return f"it's been {delay_days} days since I filed my complaint"
    elif delay_days < 180:
        return f"it's been {delay_days} days and my case is still dragging on"
    else:
        return f"it's been over {delay_days} days and nothing seems to be moving"


def derive_mood_score(row: dict, rng: random.Random) -> int:
    score = 3
    if row["case_status"] in ("Closed - insufficient evidence", "Pending review"):
        score -= 1
    if row["injury_reported"] == "Yes":
        score -= 1
    if row["delay_days"] >= 180:
        score -= 1
    if row["case_status"] == "Chargesheet filed":
        score += 1
    if row["delay_reason"] in ("Fear of retaliation", "Concern about family safety"):
        score -= 1
    score += rng.choice([-1, 0, 0, 1])  # a little natural variation
    return max(1, min(5, score))


def derive_threat_reported(row: dict, rng: random.Random) -> bool:
    if row["incident_type"] in ("Threat/intimidation", "Extortion/coercion"):
        return rng.random() < 0.65
    if row["delay_reason"] == "Fear of retaliation":
        return rng.random() < 0.5
    return rng.random() < 0.05


def derive_missed_checkins(row: dict, rng: random.Random) -> int:
    base = 0
    if row["delay_reason"] in ("Social pressure", "Concern about family safety"):
        base += rng.randint(1, 3)
    if row["follow_up_required"] == "Yes":
        base += rng.randint(0, 2)
    return min(base, 5)


def build_checkin_text(row: dict, mood_score: int, threat_reported: bool, rng: random.Random) -> str:
    parts = [delay_phrase(row["delay_days"])]

    status = STATUS_PHRASES.get(row["case_status"])
    if status:
        parts[0] += f", and {status}"

    if row["injury_reported"] == "Yes" and rng.random() < 0.6:
        parts.append(rng.choice(INJURY_CLAUSES))

    if threat_reported:
        parts.append(rng.choice(THREAT_CLAUSES))

    reason_clause = REASON_CLAUSES.get(row["delay_reason"])
    if reason_clause and rng.random() < 0.7:
        parts.append(reason_clause)

    parts.append(rng.choice(MOOD_CLAUSES[mood_score]))

    return ". ".join(p.strip().capitalize() for p in parts if p.strip()) + "."


def load_rows(xlsx_path: str) -> list[dict]:
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["Complaints"]
    headers = [c.value for c in ws[1]]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        rows.append(dict(zip(headers, r)))
    return rows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True)
    ap.add_argument("--output", default="dataset.jsonl")
    ap.add_argument("--sample-size", type=int, default=None,
                     help="Number of cases to use (default: all rows)")
    ap.add_argument("--seed", type=int, default=42)
    args = ap.parse_args()

    rng = random.Random(args.seed)
    rows = load_rows(args.input)
    if args.sample_size:
        rows = rng.sample(rows, min(args.sample_size, len(rows)))

    engine = DistressEngine()
    written = 0

    with open(args.output, "w", encoding="utf-8") as out:
        for row in rows:
            mood_score = derive_mood_score(row, rng)
            threat_reported = derive_threat_reported(row, rng)
            missed_checkins = derive_missed_checkins(row, rng)
            text = build_checkin_text(row, mood_score, threat_reported, rng)

            checkin = CheckIn(
                mood_score=mood_score,
                text=text,
                days_since_complaint=row["delay_days"] or 0,
                missed_checkins=missed_checkins,
                threat_reported=threat_reported,
            )
            result = engine.score(checkin)

            record = {
                "case_id": row["case_id"],
                "text": text,
                "severity": result.score,
                "indicators": result.reasons,
            }
            out.write(json.dumps(record) + "\n")
            written += 1

    print(f"Wrote {written} labeled examples to {args.output}")


if __name__ == "__main__":
    main()